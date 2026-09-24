#!/usr/bin/env python3
"""
Generator for Cofoundaz_QA_Test_Sheet.xlsx — the non-technical manual QA workbook.

This is the source of truth for the Excel sheet. Edit the `sections` list below
to add/change test cases, then regenerate:

    python3 -m venv .venv && . .venv/bin/activate     # or use any env
    pip install openpyxl
    python docs/checklist/build_qa_test_sheet.py

The workbook is written next to this script (docs/checklist/Cofoundaz_QA_Test_Sheet.xlsx).
Keep the IDs in sync with docs/checklist/api-test-checklist.md.
"""
import os
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.utils import get_column_letter

OUT = os.path.join(os.path.dirname(os.path.abspath(__file__)), "Cofoundaz_QA_Test_Sheet.xlsx")
GREEN="183B28"; GREEN_MID="2D5A3F"; SAND="F7EEDC"; LIGHT="F5F7F4"
ROW_A="FFFFFF"; ROW_B="EFF3F0"; GREY="617065"; RED="B0483B"; AMBER="8A6E33"; BAND="12291F"
OK_FILL="EAF2ED"; WARN_FILL="FBF3E2"; BLOCK_FILL="FBEBEB"
FONT="Arial"
thin=Side(style="thin",color="D5DDD6")
border=Border(left=thin,right=thin,top=thin,bottom=thin)

wb=Workbook()

# ============================================================ Read me first
ws=wb.active; ws.title="Read me first"; ws.sheet_view.showGridLines=False
ws.column_dimensions["A"].width=3; ws.column_dimensions["B"].width=104
def put(row,text,font=None,fill=None,wrap=True,height=None):
    c=ws.cell(row=row,column=2,value=text)
    c.font=font or Font(name=FONT,size=11,color="1E2923")
    c.alignment=Alignment(wrap_text=wrap,vertical="top")
    if fill: c.fill=PatternFill("solid",fgColor=fill)
    if height: ws.row_dimensions[row].height=height
    return c
r=2
put(r,"Cofoundaz — Manual Test Sheet",Font(name=FONT,size=18,bold=True,color=GREEN)); ws.row_dimensions[r].height=26; r+=1
put(r,"The full app walkthrough for a non-technical tester. Follow each step, compare to what should happen, and mark the result. No coding needed.",
    Font(name=FONT,size=11,italic=True,color=GREY),height=30); r+=2
put(r,"What this covers",Font(name=FONT,size=13,bold=True,color=GREEN_MID)); r+=1
put(r,"Every area of the app, grouped into 7 sections: (1) Sign-in, account & dashboard  (2) Health score  (3) Today’s Mission  "
      "(4) Roadmap  (5) Business Builder  (6) Notifications, Journal & Learning  (7) Documents. Each row is one small thing to check.",
    height=44); r+=2
put(r,"Before you start — what you’ll need",Font(name=FONT,size=13,bold=True,color=GREEN_MID)); r+=1
for item in [
    "•  The link to the test / staging web app  (ask the team for it).",
    "•  A test account login  (ask the team — we don’t put passwords in this sheet).",
    "•  A SPARE / throwaway account for the two-step-verification tests (1.9 & 1.10) — see the warning below.",
    "•  A phone with a free authenticator app  (Google Authenticator or Authy) for those two tests.",
    "•  Ideally a second, lower-scoring account to fully test Health recommendations (2.5).",
]:
    put(r,item,height=26); r+=1
r+=1
put(r,"⚠️  IMPORTANT (for tests 1.9 & 1.10):  Once two-step verification is turned ON there is currently NO way to turn it off. "
      "Do NOT enable it on the shared / main test account — use a throwaway account only.",
    Font(name=FONT,size=11,bold=True,color="7A2E22"),fill=SAND,height=44); r+=2
put(r,"How to use this sheet",Font(name=FONT,size=13,bold=True,color=GREEN_MID)); r+=1
for item in [
    "1.  Go to the “Test cases” tab and work down the list (it’s grouped by section).",
    "2.  For each row: do the steps in “What to do”, then compare what you see to “What should happen”.",
    "3.  The “Expected” column tells you upfront whether it should work, work-with-a-catch, or isn’t ready yet.",
    "4.  Read “Heads-up” — it warns you about things that look odd but are normal.",
    "5.  Set the “Result” dropdown and add anything useful in “Your notes” (a screenshot link helps a lot for a Fail).",
    "6.  Check the “Known — not ready yet” tab BEFORE logging a bug.",
]:
    put(r,item,height=24); r+=1
r+=1
put(r,"The “Expected” column",Font(name=FONT,size=13,bold=True,color=GREEN_MID)); r+=1
for txt,fill in [("✅  Works — should behave exactly as described.",OK_FILL),
                 ("⚠️  Works, but watch out — there’s a catch noted in “Heads-up”.",WARN_FILL),
                 ("⛔  Not ready yet — expected to be unfinished (backend). Don’t log as a bug.",BLOCK_FILL)]:
    put(r,txt,fill=fill,height=24); r+=1
r+=1
put(r,"The “Result” options",Font(name=FONT,size=13,bold=True,color=GREEN_MID)); r+=1
for txt in ["Pass  —  worked as “What should happen” describes.",
            "Fail  —  did something different or broke (add details + a screenshot).",
            "Blocked  —  couldn’t test it (say why).",
            "Not tested  —  skipped / didn’t get to it."]:
    put(r,txt,height=22); r+=1
r+=1
put(r,"Something clearly broken that isn’t in “Known — not ready yet”? Flag it to the dev team with the test ID (e.g. 4.7) and a screenshot.",
    Font(name=FONT,size=10,italic=True,color=GREY),height=30)

# ============================================================ Test cases
tc=wb.create_sheet("Test cases"); tc.sheet_view.showGridLines=False
headers=["ID","What we're testing","Where to go","What to do","What should happen","Heads-up (this is normal)","Expected","Result","Your notes"]
widths =[7,   24,                   24,           44,           40,                   38,                         16,        13,        26]

# sections: (title, [ (id, what, where, do, happen, headsup, flag) ])
sections=[
 ("1 · Sign-in, account & dashboard", [
   ("1.1","Log in","Sign-in page","Log in with a valid email and password.","You land on your dashboard and stay signed in.","A wrong password shows an inline error — it should NOT loop you around.","ok"),
   ("1.2","Stay signed in when idle","Any page","Leave the tab open a long time (past a normal timeout), then click around or do an action.","Your action still works — you’re not kicked out.","You should not get logged out or see errors just from being idle.","ok"),
   ("1.3","Expired login is handled","Any page","If your login truly expires, do any action. (Edge case — skip if you can’t force it.)","You’re taken to the sign-in page with a “your session expired” note — once, not repeatedly.","This is hard to trigger on purpose; mark Not tested if you can’t.","ok"),
   ("1.4","Email links work","From an email","Open a “verify your email” or “reset password” link.","The link opens the correct page and works.","—","ok"),
   ("1.5","Sign-up steps resume","Onboarding","Start the sign-up steps, fill in a couple, then refresh the page.","It picks up where you left off.","Pick your country from the dropdown (you can’t free-type a country).","ok"),
   ("1.6","Upload your company logo","Onboarding → logo step","Choose and upload a logo image.","The logo uploads and a preview appears.","It now saves to the real system — confirm the preview actually shows your image.","warn"),
   ("1.7","Dashboard shows real numbers","Dashboard (home)","Open the dashboard.","The cards, the health-score pill and the notification bell show real numbers for your account.","No fixed placeholder numbers.","ok"),
   ("1.8","“Do it” on the AI briefing card","Dashboard","Click the “Do it” button on the AI briefing card.","The action is accepted and the page does not crash.","If that action isn’t switched on in the backend yet it may quietly do nothing — expected, not a crash.","warn"),
   ("1.9","Turn on two-step verification","Left sidebar → “Security (2FA)”","On a THROWAWAY account: open an authenticator app on your phone, scan the QR (or type the key), enter the 6-digit code, then save the backup codes shown.","Two-step verification switches on and a set of one-time backup codes appears to copy or download.","⚠️ Once ON it can’t be turned off yet — throwaway account only. “SMS text message” is greyed out on purpose.","ok"),
   ("1.10","Sign in with two-step verification","Sign-in page","On that account: sign out, sign back in, and enter the 6-digit code when asked. Also try “Use a backup code” with a saved code.","After the password you’re asked for a 6-digit code; a correct code (or backup code) signs you in.","Opening the code screen directly without signing in first should send you back to the sign-in page.","ok"),
   ("1.11","Accept a team invitation","The invite link someone sends you","Open the invite link — try it once signed OUT and once signed IN.","You see who invited you, the workspace and your role. Signed in → “Accept” adds you and opens the dashboard. Signed out → log-in / sign-up buttons that bring you back to accept.","Signed in with a DIFFERENT email than invited → a message to log in with the invited email. An old/used link shows “no longer valid”.","ok"),
 ]),
 ("2 · Health score", [
   ("2.1","Overview loads","Health","Open Health.","You see your overall score and the five areas.","A thriving account shows around 90.","ok"),
   ("2.2","Open the Financial area","Health → Dimensions → Financial","Open the “Financial” area.","It loads with its details.","Make sure “Financial” isn’t blank (it’s stored under a different name internally).","warn"),
   ("2.3","Change the history range","Health → Trend history","Switch the time range (7d / 30d / 90d / all).","The chart updates to match the range you picked.","—","ok"),
   ("2.4","Benchmarks","Health → Benchmarks","Open Benchmarks.","You see comparison data, OR an honest “not enough data yet” message.","If empty it should SAY so — not show made-up numbers.","warn"),
   ("2.5","Accept / dismiss a recommendation","Health → Recommendations","Accept or dismiss a recommendation.","It resolves and leaves the list.","The main test account may have zero — try a lower-scoring account. Acting on the same one twice may say “already resolved”.","warn"),
   ("2.6","Charts show real data","Health → overview, Trend history, a single dimension","Look at the three charts: the “Dimension balance” wheel, the “Score over time” chart, and a dimension’s trend chart.","The wheel shows your five scores; the score-over-time chart has a 0–100 scale and a tooltip; a dimension shows its own trend.","Only one assessment done → the score chart shows a single dot (correct). The dashboard KPI cards no longer show a mini trend line (there’s no real trend data for them).","ok"),
 ]),
 ("3 · Today’s Mission", [
   ("3.1","Load today’s mission","Today’s Mission","Open it.","You see today’s tasks and your streak.","A weekend-off day can legitimately have no tasks.","ok"),
   ("3.2","Task actions","Today’s Mission","Complete, snooze, reject, add, or reorder a task.","Each updates, and the list and streak refresh correctly.","Completing a task is final — you can’t un-complete it.","ok"),
   ("3.3","Reject a task","Today’s Mission","Reject a task and pick a reason.","It’s rejected with the reason you chose.","Use the reason chips provided (don’t type your own).","warn"),
   ("3.4","Mission settings","Mission → Settings","Change mission size, delivery time, or the weekends option.","Your settings save.","Time flips between 12-hour and 24-hour; “weekends off” is the opposite of “weekend missions”.","ok"),
   ("3.5","History pages","Mission → Completed / Streaks / Upcoming","Open each.","They fill with your real history.","No history yet → an honest empty message.","ok"),
 ]),
 ("4 · Roadmap", [
   ("4.1","Load the roadmap","Roadmap","Open it.","You see phases → milestones → tasks.","A “slipped / behind” count shows if any dates have slipped.","ok"),
   ("4.2","Add / edit / delete","Roadmap","Add, edit, or delete a phase, milestone or task.","The change saves and the roadmap updates.","Editors only — a mentor/viewer can’t make changes.","ok"),
   ("4.3","Block a dependency loop","Roadmap → Dependencies","Try to link two tasks so they depend on each other in a circle.","The app refuses to create the circular link.","It should block the loop.","warn"),
   ("4.4","Apply a template","Roadmap → Templates","Browse the gallery and apply a template.","The roadmap updates from the template.","—","ok"),
   ("4.5","AI Re-plan — preview & apply","Roadmap → Re-Plan","Preview an AI re-plan, then apply it, and check history.","Preview shows changes; applying updates the plan; history logs it.","Nothing changes until you click Apply.","ok"),
   ("4.6","Move a task on the Kanban board","Roadmap → Kanban (Status view)","Drag a task card between columns (To Do / In Progress / Done).","The card moves and the change stays after a refresh.","Mentor/viewer → the card snaps back (permission working). The “Phase” view isn’t drag-and-drop; tap a card to open it.","ok"),
   ("4.7","Re-plan — date changes & history","Roadmap → Re-Plan","Generate a re-plan; expand a past entry under “Re-plan History”.","Each change shows old date → new date and how far it moved (e.g. “+17 days”); history entries expand to their before/after dates.","Deltas come from the real dates — nothing invented.","ok"),
 ]),
 ("5 · Business Builder", [
   ("5.1","Overview","Business Builder","Open it.","You see each canvas with a completion %.","—","ok"),
   ("5.2","Lean canvas saves","Business Builder → Lean Canvas","Edit some blocks and wait a moment.","It saves on its own.","Advanced: editing the same canvas in two tabs may briefly show it reconciling — that’s the safety check.","warn"),
   ("5.3","Canvases save automatically","Business Model / Value Proposition / SWOT / Mission & Vision","Edit and pause.","It saves on its own after you stop typing.","Pause after typing to let it save.","ok"),
   ("5.4","Records — add / edit / delete","Personas / Competitive / Pricing / Revenue","Create, edit, and delete a record.","The list updates.","Leaving a required field empty shows a validation message. Pricing has tiers inside it.","warn"),
   ("5.5","Suggestions","Business Builder → Suggestions","Approve or reject a suggestion.","It applies or updates.","May say “no longer pending” if the canvas has moved on.","warn"),
   ("5.6","Positioning map","Business Builder → Positioning map","Edit the axes.","Competitors plot against your saved axes.","—","ok"),
   ("5.7","AI-fill / AI draft","Any canvas","Click AI-fill / AI draft.","It says “queued — coming soon”.","It never actually finishes yet — expected, not a bug.","blocked"),
   ("5.8","Edit & reorder canvas notes","Business Model / Value Proposition / SWOT","Click a note to edit it in place (Enter saves, Esc cancels). Then hover a note, grab its drag handle, and drag it to reorder within the same box.","The note updates and saves; the reordered notes save in the new order.","Emptying a note and pressing Enter deletes it. Reorder stays within one box. “Mission & Vision” is a text box; “Lean Canvas” doesn’t have these yet.","ok"),
 ]),
 ("6 · Notifications, Journal & Learning", [
   ("6.1","Notifications inbox","Bell / Notifications","Open it, mark one read, and “mark all read”.","The unread count updates.","—","ok"),
   ("6.2","Email preferences","Notifications → Preferences","Toggle the master email switch and the 5 category switches.","They save.","The “Digest & quiet hours” and “Announcements” tabs are demo-only for now.","ok"),
   ("6.3","Read the journal","Journal","Read journal entries.","Reading works.","—","ok"),
   ("6.4","Save a new journal entry","Journal","Try to save a new entry.","It fails gracefully — it doesn’t crash the page.","Saving isn’t set up on the test server yet — expected to fail softly.","blocked"),
   ("6.5","Learning academy","Academy","Browse courses / articles / paths / certificates, enroll, and complete a lesson.","These work and save.","—","ok"),
 ]),
 ("7 · Documents", [
   ("7.1","Files","Documents","Upload, list, and delete a file.","All three work.","—","ok"),
   ("7.2","Share a document","Documents","Share a document, then remove the share.","The share list updates.","—","ok"),
   ("7.3","Signature requests","Documents → Signatures","Send, remind, or cancel a signature request.","The status reflects what you did.","—","ok"),
   ("7.4","Create from a template","Documents → Templates","Create a document from a template.","It creates the document.","—","ok"),
   ("7.5","Edit & save a document","Documents → open a document","Open the editor, make edits, and save.","Your edits save.","Documents don’t auto-save — click Save. If it was changed elsewhere you get a “reload” banner and it won’t overwrite.","warn"),
   ("7.6","Open a real emailed sign/share link","The sign/share link","Open a real emailed link.","The recipient page works when opened on the app website.","Straight from an email it may show a raw data page — open it on the app website instead.","blocked"),
   ("7.7","Edit a document’s sections","Documents → open a document","Add a section, delete one, drag a section to reorder (or use the up/down arrows), and rename a heading. Then Save.","All the section changes save after you click Save.","Documents don’t auto-save — you must click Save. Headings used to be read-only.","ok"),
   ("7.8","Sign a document","The signing link, on the app website","Read the preview, type your full name (shown styled like a signature), tick the “I agree…” box, then Sign.","Your typed name is accepted and you see a “signed — thank you” confirmation.","It’s a typed signature (no drawing pad). The Sign button stays disabled until you tick consent.","ok"),
 ]),
]

# title band
tc.merge_cells("A1:I1")
t=tc.cell(row=1,column=1,value="Test cases — full app walkthrough")
t.font=Font(name=FONT,size=14,bold=True,color="FFFFFF"); t.alignment=Alignment(vertical="center",indent=1)
t.fill=PatternFill("solid",fgColor=GREEN); tc.row_dimensions[1].height=26
# header
for i,(h,w) in enumerate(zip(headers,widths),start=1):
    c=tc.cell(row=2,column=i,value=h)
    c.font=Font(name=FONT,size=10,bold=True,color="FFFFFF")
    c.fill=PatternFill("solid",fgColor=GREEN_MID)
    c.alignment=Alignment(wrap_text=True,vertical="center",horizontal="center")
    c.border=border
    tc.column_dimensions[get_column_letter(i)].width=w
tc.row_dimensions[2].height=30

EXP={"ok":("✅  Works",OK_FILL,GREEN_MID),
     "warn":("⚠️  Works, watch out",WARN_FILL,AMBER),
     "blocked":("⛔  Not ready (expected)",BLOCK_FILL,RED)}

rr=3; band_toggle=0
dv=DataValidation(type="list",formula1='"Pass,Fail,Blocked,Not tested"',allow_blank=True)
dv.promptTitle="Result"; dv.prompt="Pick one"; tc.add_data_validation(dv)
first_data=None; last_data=None
for title,items in sections:
    # section band
    tc.merge_cells(start_row=rr,start_column=1,end_row=rr,end_column=9)
    b=tc.cell(row=rr,column=1,value=title)
    b.font=Font(name=FONT,size=11,bold=True,color="FFFFFF")
    b.alignment=Alignment(vertical="center",indent=1)
    b.fill=PatternFill("solid",fgColor=BAND)
    tc.row_dimensions[rr].height=22
    rr+=1
    for idx,(cid,what,where,do,happen,heads,flag) in enumerate(items):
        fill=ROW_A if idx%2==0 else ROW_B
        exp_label,exp_fill,exp_txt=EXP[flag]
        vals=[cid,what,where,do,happen,heads,exp_label,"",""]
        for col in range(1,10):
            c=tc.cell(row=rr,column=col,value=vals[col-1])
            c.border=border
            c.alignment=Alignment(wrap_text=True,vertical="top",horizontal="center" if col in(1,7,8) else "left")
            if col==7:
                c.fill=PatternFill("solid",fgColor=exp_fill)
                c.font=Font(name=FONT,size=9,bold=True,color=exp_txt)
            else:
                c.fill=PatternFill("solid",fgColor=fill)
                if col==1: c.font=Font(name=FONT,size=10,bold=True,color=GREEN_MID)
                elif col==2: c.font=Font(name=FONT,size=10,bold=True,color="1E2923")
                elif col==6: c.font=Font(name=FONT,size=9,italic=True,color=AMBER)
                else: c.font=Font(name=FONT,size=10,color="1E2923")
        tc.row_dimensions[rr].height=92
        if first_data is None: first_data=rr
        last_data=rr
        rr+=1
dv.add(f"H{first_data}:H{last_data}")
tc.freeze_panes="A3"

# ============================================================ Known - not ready yet
kn=wb.create_sheet("Known - not ready yet"); kn.sheet_view.showGridLines=False
kn.merge_cells("A1:B1")
kt=kn.cell(row=1,column=1,value="Known — NOT ready yet  (please don’t log these as bugs)")
kt.font=Font(name=FONT,size=14,bold=True,color="FFFFFF"); kt.alignment=Alignment(vertical="center",indent=1)
kt.fill=PatternFill("solid",fgColor=AMBER); kn.row_dimensions[1].height=26
kn.column_dimensions["A"].width=34; kn.column_dimensions["B"].width=88
known=[
 ("Assessment questions","The assessment pages are built, but the question wording isn’t coming from the backend yet — so real questions can’t show."),
 ("“AI draft” / “AI-fill” on canvases","Clicking it says “queued — coming soon” and never finishes. There’s no AI worker running yet. It’s an honest placeholder (test 5.7)."),
 ("Saving a new Journal entry","Reading the journal works, but SAVING a new entry fails on the test server (an encryption key isn’t set there yet) (test 6.4)."),
 ("Sign / share links from an email","A link opened straight from an email may show a plain data page. Open the link on the app website instead (tests 7.6, 7.8)."),
 ("Two-step verification by SMS","The “SMS text message” option is greyed out (“coming soon”). Only the authenticator-app option works."),
 ("Turning OFF two-step verification","There’s no “turn off” yet — that’s why tests 1.9 / 1.10 must use a throwaway account."),
 ("Demo-only areas (placeholder data)","Not connected to real data yet — don’t test as real: Marketing, Sales, Finance, Validation, Funding, Investor Readiness, Legal & Compliance, Calendar, Analytics & Reports, Marketplace, Admin / Super-Admin, and the Notifications “Digest / quiet hours” & “Announcements” tabs."),
]
kr=3
for a,d in known:
    ca=kn.cell(row=kr,column=1,value=a); ca.font=Font(name=FONT,size=11,bold=True,color=GREEN_MID)
    ca.alignment=Alignment(wrap_text=True,vertical="top"); ca.fill=PatternFill("solid",fgColor=LIGHT); ca.border=border
    cb=kn.cell(row=kr,column=2,value=d); cb.font=Font(name=FONT,size=11,color="1E2923")
    cb.alignment=Alignment(wrap_text=True,vertical="top"); cb.fill=PatternFill("solid",fgColor=LIGHT); cb.border=border
    kn.row_dimensions[kr].height=50; kr+=1
kn.freeze_panes="A3"

wb.save(OUT)
print("saved",OUT)
print("total test cases:", sum(len(items) for _,items in sections))
